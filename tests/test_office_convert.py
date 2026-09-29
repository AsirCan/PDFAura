"""Issue #7: Office conversions need CoInitialize and must not touch the user's Office.

There is no Office on the test machine, so a fake pythoncom/win32com.client
pair is injected into sys.modules and the COM calls are recorded.
"""
import sys
import types

import pytest

from src.core import convert


class FakeComError(Exception):
    def __init__(self, hresult):
        super().__init__(hresult, "fake com error")
        self.hresult = hresult


class Recorder:
    def __init__(self):
        self.events = []

    def log(self, name, *args):
        self.events.append((name,) + args)

    def names(self):
        return [e[0] for e in self.events]


class FakeDocument:
    def __init__(self, rec, kind, open_kwargs):
        self.rec = rec
        self.kind = kind
        self.open_kwargs = open_kwargs

    def ExportAsFixedFormat(self, *args, **kwargs):
        self.rec.log("export", self.kind, args)

    def SaveAs(self, *args):
        self.rec.log("saveas", self.kind, args)

    def Close(self, *args):
        self.rec.log("close", self.kind)


class FakeCollection:
    def __init__(self, rec, kind, count=0):
        self.rec = rec
        self.kind = kind
        self.Count = count

    def Open(self, path, **kwargs):
        self.rec.log("open", self.kind, kwargs)
        return FakeDocument(self.rec, self.kind, kwargs)


class FakeApp:
    def __init__(self, rec, prog_id):
        self.rec = rec
        self.prog_id = prog_id
        self.Documents = FakeCollection(rec, "word")
        self.Presentations = FakeCollection(rec, "ppt")
        self.Workbooks = FakeCollection(rec, "excel")

    def __setattr__(self, name, value):
        if name in ("DisplayAlerts", "Visible", "Interactive"):
            self.rec.log("set", name, value)
        object.__setattr__(self, name, value)

    def Quit(self):
        self.rec.log("quit", self.prog_id)


@pytest.fixture
def fake_com(monkeypatch):
    """Install a fake pythoncom + win32com.client and record what is called."""
    rec = Recorder()
    rec.running = set()          # prog_ids the "user" already has open
    rec.dispatch_error = None    # HRESULT to raise from DispatchEx

    pythoncom = types.ModuleType("pythoncom")
    pythoncom.com_error = FakeComError
    pythoncom.CoInitialize = lambda: rec.log("CoInitialize")
    pythoncom.CoUninitialize = lambda: rec.log("CoUninitialize")

    def dispatch_ex(prog_id):
        rec.log("DispatchEx", prog_id)
        if rec.dispatch_error is not None:
            raise FakeComError(rec.dispatch_error)
        return FakeApp(rec, prog_id)

    def get_active_object(prog_id):
        if prog_id in rec.running:
            return FakeApp(rec, prog_id)
        raise FakeComError(-2147221021)

    client = types.ModuleType("win32com.client")
    client.DispatchEx = dispatch_ex
    client.Dispatch = lambda prog_id: pytest.fail("Dispatch attaches to the user's Office")
    client.GetActiveObject = get_active_object

    win32com = types.ModuleType("win32com")
    win32com.client = client

    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)
    monkeypatch.setitem(sys.modules, "win32com", win32com)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    return rec


CONVERTERS = [
    (convert.word_to_pdf, "in.docx", "Word.Application"),
    (convert.ppt_to_pdf, "in.pptx", "PowerPoint.Application"),
    (convert.excel_to_pdf, "in.xlsx", "Excel.Application"),
]


@pytest.mark.parametrize("func, name, prog_id", CONVERTERS)
def test_conversion_initialises_com_on_this_thread(tmp_path, fake_com, func, name, prog_id):
    """Without CoInitialize every conversion fails with 'CoInitialize was not called'."""
    src = tmp_path / name
    src.write_bytes(b"x")

    func(str(src), str(tmp_path / "out.pdf"))

    assert fake_com.names()[0] == "CoInitialize"
    assert fake_com.names()[-1] == "CoUninitialize"


@pytest.mark.parametrize("func, name, prog_id", CONVERTERS)
def test_conversion_uses_a_private_instance(tmp_path, fake_com, func, name, prog_id):
    """Dispatch would attach to the user's Office; DispatchEx asks for our own."""
    src = tmp_path / name
    src.write_bytes(b"x")

    func(str(src), str(tmp_path / "out.pdf"))

    assert ("DispatchEx", prog_id) in fake_com.events


@pytest.mark.parametrize("func, name, prog_id", CONVERTERS)
def test_documents_are_opened_read_only(tmp_path, fake_com, func, name, prog_id):
    src = tmp_path / name
    src.write_bytes(b"x")

    func(str(src), str(tmp_path / "out.pdf"))

    opens = [e for e in fake_com.events if e[0] == "open"]
    assert opens and opens[0][2].get("ReadOnly") is True


@pytest.mark.parametrize("func, name, prog_id", CONVERTERS)
def test_our_own_document_is_closed(tmp_path, fake_com, func, name, prog_id):
    src = tmp_path / name
    src.write_bytes(b"x")

    func(str(src), str(tmp_path / "out.pdf"))

    assert "close" in fake_com.names()


def test_powerpoint_already_open_is_not_quit(tmp_path, fake_com):
    """PowerPoint is single-instance: quitting it closes the user's decks."""
    fake_com.running.add("PowerPoint.Application")
    src = tmp_path / "in.pptx"
    src.write_bytes(b"x")

    convert.ppt_to_pdf(str(src), str(tmp_path / "out.pdf"))

    assert "quit" not in fake_com.names()


def test_powerpoint_we_started_is_quit(tmp_path, fake_com):
    src = tmp_path / "in.pptx"
    src.write_bytes(b"x")

    convert.ppt_to_pdf(str(src), str(tmp_path / "out.pdf"))

    assert ("quit", "PowerPoint.Application") in fake_com.events


@pytest.fixture
def no_libreoffice(monkeypatch):
    from src.utils import libreoffice_helper
    monkeypatch.setattr(libreoffice_helper, "find_libreoffice", lambda: None)


@pytest.mark.parametrize("func, name, prog_id", CONVERTERS)
def test_missing_office_gives_readable_message(tmp_path, fake_com, no_libreoffice, func, name, prog_id):
    fake_com.dispatch_error = -2147221005  # invalid class string
    src = tmp_path / name
    src.write_bytes(b"x")

    with pytest.raises(RuntimeError) as excinfo:
        func(str(src), str(tmp_path / "out.pdf"))

    message = str(excinfo.value)
    assert "-2147221005" not in message
    assert prog_id.split(".")[0] in message
    # The thread is still uninitialised even though dispatch failed.
    assert fake_com.names()[-1] == "CoUninitialize"


@pytest.mark.parametrize("func, name, prog_id", CONVERTERS)
def test_failure_still_uninitialises_and_quits(tmp_path, fake_com, func, name, prog_id):
    src = tmp_path / name
    src.write_bytes(b"x")

    def boom(*a, **k):
        raise FakeComError(-2147352567)

    monkey = {"Word.Application": "Documents", "PowerPoint.Application": "Presentations",
              "Excel.Application": "Workbooks"}[prog_id]
    original = FakeCollection.Open
    FakeCollection.Open = boom
    try:
        with pytest.raises(RuntimeError):
            func(str(src), str(tmp_path / "out.pdf"))
    finally:
        FakeCollection.Open = original

    assert monkey  # sanity: the collection name is the one we expect to fail
    assert fake_com.names()[-1] == "CoUninitialize"


# ── LibreOffice fallback ───────────────────────────────────────────────────
#
# PowerPoint → PDF failed with "PowerPoint is not installed" on a machine
# that had LibreOffice, which converts the same files.


class FakeSoffice:
    """Stands in for subprocess.Popen running soffice --convert-to pdf."""

    def __init__(self, rec, write_output=True, returncode=0, hang=False):
        self.rec = rec
        self.write_output = write_output
        self.returncode_on_exit = returncode
        self.hang = hang

    def __call__(self, command, **kwargs):
        self.rec.log("soffice", command, kwargs)
        out_dir = command[command.index("--outdir") + 1]
        if self.write_output and not self.hang:
            with open(f"{out_dir}/input.pdf", "wb") as f:
                f.write(b"%PDF-1.7 from libreoffice")
        return FakeProcess(self)


class FakeProcess:
    pid = 4242

    def __init__(self, fake):
        self.fake = fake
        self.returncode = None if fake.hang else fake.returncode_on_exit

    def communicate(self, timeout=None):
        import subprocess
        if self.returncode is None:
            raise subprocess.TimeoutExpired("soffice", timeout)
        return b"soffice output", None

    def poll(self):
        return self.returncode

    def kill(self):
        self.fake.rec.log("kill")
        self.returncode = -9


@pytest.fixture
def libreoffice(monkeypatch, fake_com):
    """Office is missing, LibreOffice is 'installed' and 'converts' the file."""
    import subprocess
    from src.utils import libreoffice_helper

    fake_com.dispatch_error = -2147221005
    monkeypatch.setattr(libreoffice_helper, "find_libreoffice", lambda: r"C:\LO\soffice.exe")
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: fake_com.log("taskkill"))
    fake = FakeSoffice(fake_com)
    monkeypatch.setattr(subprocess, "Popen", fake)
    return fake


@pytest.mark.parametrize("func, name, prog_id", CONVERTERS)
def test_libreoffice_is_used_when_office_is_missing(tmp_path, libreoffice, func, name, prog_id):
    src = tmp_path / name
    src.write_bytes(b"x")
    out = tmp_path / "out.pdf"

    func(str(src), str(out))

    assert out.read_bytes() == b"%PDF-1.7 from libreoffice"
    command = [e for e in libreoffice.rec.events if e[0] == "soffice"][0][1]
    assert command[0] == r"C:\LO\soffice.exe"
    assert "--headless" in command
    assert command[command.index("--convert-to") + 1] == "pdf"


def test_libreoffice_gets_a_private_profile(tmp_path, libreoffice):
    """Sharing the user's profile hands the job to their open LibreOffice,
    which returns at once without writing anything."""
    src = tmp_path / "in.pptx"
    src.write_bytes(b"x")

    convert.ppt_to_pdf(str(src), str(tmp_path / "out.pdf"))

    command = [e for e in libreoffice.rec.events if e[0] == "soffice"][0][1]
    profile = [arg for arg in command if arg.startswith("-env:UserInstallation=file:")]
    assert profile


def test_libreoffice_does_not_touch_the_input_folder(tmp_path, libreoffice):
    """soffice writes lock files next to what it opens."""
    folder = tmp_path / "Ders Notları"
    folder.mkdir()
    src = folder / "ders 1.pptx"
    src.write_bytes(b"x")

    convert.ppt_to_pdf(str(src), str(folder / "ders 1.pdf"))

    assert sorted(p.name for p in folder.iterdir()) == ["ders 1.pdf", "ders 1.pptx"]
    command = [e for e in libreoffice.rec.events if e[0] == "soffice"][0][1]
    assert str(src) not in command


def test_libreoffice_failure_is_readable_and_keeps_the_old_output(tmp_path, libreoffice):
    libreoffice.write_output = False
    libreoffice.returncode_on_exit = 1
    src = tmp_path / "in.pptx"
    src.write_bytes(b"x")
    out = tmp_path / "out.pdf"
    out.write_bytes(b"previous")

    with pytest.raises(RuntimeError) as excinfo:
        convert.ppt_to_pdf(str(src), str(out))

    assert "LibreOffice" in str(excinfo.value)
    assert out.read_bytes() == b"previous"


def test_cancel_stops_libreoffice(tmp_path, libreoffice):
    from src.core.task_manager import CancelledError, TaskContext

    libreoffice.hang = True
    src = tmp_path / "in.pptx"
    src.write_bytes(b"x")
    ctx = TaskContext()
    # Cancel as soon as the conversion has started.
    ctx.report_progress = lambda *a, **k: ctx.cancel() if a[0] == 0 else None

    with pytest.raises(CancelledError):
        convert.ppt_to_pdf(str(src), str(tmp_path / "out.pdf"), ctx=ctx)

    assert "taskkill" in libreoffice.rec.names() or "kill" in libreoffice.rec.names()
    assert not (tmp_path / "out.pdf").exists()


def test_office_is_preferred_when_installed(tmp_path, fake_com, monkeypatch):
    from src.utils import libreoffice_helper
    monkeypatch.setattr(libreoffice_helper, "find_libreoffice",
                        lambda: pytest.fail("LibreOffice used although Office is there"))
    src = tmp_path / "in.pptx"
    src.write_bytes(b"x")

    convert.ppt_to_pdf(str(src), str(tmp_path / "out.pdf"))

    assert ("saveas", "ppt", (str(tmp_path / "out.pdf"), convert.PP_SAVE_AS_PDF)) in fake_com.events


@pytest.mark.skipif(not __import__("src.utils.libreoffice_helper", fromlist=["x"]).find_libreoffice(),
                    reason="LibreOffice is not installed")
def test_real_libreoffice_converts_a_word_document(tmp_path, monkeypatch):
    """End to end with the real soffice, when this machine has it."""
    import docx
    import fitz

    monkeypatch.setattr(convert, "_word_to_pdf_com", _raise_not_installed)
    document = docx.Document()
    document.add_paragraph("PDF Aura LibreOffice çıktısı")
    src = tmp_path / "Ders Notları" / "ders.docx"
    src.parent.mkdir()
    document.save(src)
    out = tmp_path / "Ders Notları" / "ders.pdf"

    convert.word_to_pdf(str(src), str(out))

    with fitz.open(out) as pdf:
        assert "PDF Aura LibreOffice" in pdf[0].get_text()


def _raise_not_installed(*_args):
    raise convert.OfficeNotInstalledError("not installed")


def test_word_no_longer_uses_docx2pdf():
    """docx2pdf never calls CoInitialize and quits the user's Word."""
    import inspect
    source = inspect.getsource(convert)
    assert "docx2pdf" not in source
